#!/usr/bin/env python3
"""Génère une feuille de validation Excel à partir d'un JSON de propositions.

Usage:
    python make_review.py propositions.json [--out reviews/review-<date>.xlsx]

Le JSON attendu :
{
  "session": "Grooming backlog — équipement",
  "skill": "lb-grooming",
  "propositions": [
    {
      "ref": "P1",
      "type": "NOUVEAU TICKET",      # ou MODIF TICKET / FUSION / SUPPRESSION /
                                     # REPRIORISATION / STATU QUO / ADR / OPTION DESIGN
      "cible": "EQSET02",            # ticket/ADR concerné, ou "-" si nouveau
      "titre": "Ajouter bonus de set 2 pièces",
      "detail": "…",                 # ce qui serait écrit, concrètement
      "justification": "…",          # pourquoi
      "cout": "M",                   # S / M / L / -
      "reco": "OUI"                  # la reco de Claude : OUI / NON / A DISCUTER
    }
  ]
}

Adrian remplit la colonne DÉCISION (OUI / NON / PLUS TARD / MODIFIER) + Commentaire,
sauvegarde, puis `read_review.py` relit le fichier.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DECISIONS = ["OUI", "NON", "PLUS TARD", "MODIFIER"]

HEADERS = [
    ("Réf", 8),
    ("Type", 18),
    ("Cible", 14),
    ("Titre", 38),
    ("Détail de la proposition", 60),
    ("Justification", 40),
    ("Coût", 7),
    ("Reco Claude", 12),
    ("✅ DÉCISION", 14),
    ("Commentaire / correction", 45),
]

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # jaune = à remplir par Adrian
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(data: dict, out_path: Path) -> Path:
    props = data.get("propositions", [])
    if not props:
        sys.exit("Aucune proposition dans le JSON.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Validation"

    # --- Bandeau ---
    ws["A1"] = data.get("session", "Session de validation")
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = (
        f"Skill: {data.get('skill', '-')}  ·  Date: {date.today():%Y-%m-%d}  ·  "
        f"{len(props)} proposition(s)"
    )
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")
    ws["A3"] = (
        "LÉGENDE — Remplis uniquement les colonnes jaunes. DÉCISION : OUI (applique) · "
        "NON (abandonne) · PLUS TARD (parking, réouvrable) · MODIFIER (précise en commentaire). "
        "Une ligne laissée vide = NON TRAITÉE, Claude ne l'appliquera pas."
    )
    ws["A3"].font = Font(name=FONT, size=9, color="C00000")

    header_row = 5

    # --- En-têtes ---
    for col, (label, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width

    # --- Lignes ---
    for i, p in enumerate(props):
        r = header_row + 1 + i
        values = [
            p.get("ref", f"P{i + 1}"),
            p.get("type", ""),
            p.get("cible", "-"),
            p.get("titre", ""),
            p.get("detail", ""),
            p.get("justification", ""),
            p.get("cout", "-"),
            p.get("reco", ""),
            "",  # DÉCISION — à remplir
            "",  # Commentaire — à remplir
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if col in (9, 10):  # colonnes de saisie
                cell.fill = INPUT_FILL
        ws.row_dimensions[r].height = 58

    first, last = header_row + 1, header_row + len(props)

    # --- Liste déroulante sur DÉCISION ---
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(DECISIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "Valeurs autorisées : OUI, NON, PLUS TARD, MODIFIER"
    dv.errorTitle = "Décision invalide"
    ws.add_data_validation(dv)
    dv.add(f"I{first}:I{last}")

    # --- Compteur de suivi (formules Excel-2007-safe) ---
    s = last + 2
    ws.cell(row=s, column=8, value="Traitées :").font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=s, column=9, value=f'=COUNTA(I{first}:I{last})&" / {len(props)}"').font = Font(
        name=FONT, size=10, bold=True
    )
    ws.cell(row=s + 1, column=8, value="dont OUI :").font = Font(name=FONT, size=10)
    ws.cell(row=s + 1, column=9, value=f'=COUNTIF(I{first}:I{last},"OUI")').font = Font(
        name=FONT, size=10
    )

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:J{last}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("json_path", type=Path)
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()

    data = json.loads(args.json_path.read_text(encoding="utf-8"))
    out = args.out or Path("reviews") / f"review-{date.today():%Y%m%d}.xlsx"
    saved = build(data, out)
    print(f"✅ Feuille de validation : {saved}")
    print("   Remplis la colonne DÉCISION, sauvegarde, puis : python read_review.py", saved)


if __name__ == "__main__":
    main()
