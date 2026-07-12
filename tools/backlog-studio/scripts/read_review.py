#!/usr/bin/env python3
"""Relit une feuille de validation remplie et sort les décisions en JSON.

Usage:
    python read_review.py reviews/review-20260711.xlsx

Sortie (stdout, JSON) — Claude ne doit appliquer QUE le bloc "appliquer".
Toute ligne sans décision est listée dans "non_traitees" et NE DOIT PAS être appliquée.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

VALID = {"OUI", "NON", "PLUS TARD", "MODIFIER"}
HEADER_ROW = 5
COLS = {
    "ref": 1, "type": 2, "cible": 3, "titre": 4, "detail": 5,
    "justification": 6, "cout": 7, "reco": 8, "decision": 9, "commentaire": 10,
}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx_path", type=Path)
    args = ap.parse_args()

    if not args.xlsx_path.exists():
        sys.exit(f"Fichier introuvable : {args.xlsx_path}")

    ws = load_workbook(args.xlsx_path, data_only=True).active

    buckets: dict[str, list] = {
        "appliquer": [], "abandonner": [], "parking": [],
        "a_modifier": [], "non_traitees": [], "invalides": [],
    }

    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        ref = row[COLS["ref"] - 1].value
        if not ref:
            continue
        item = {
            k: (row[i - 1].value or "") for k, i in COLS.items() if k != "decision"
        }
        raw = row[COLS["decision"] - 1].value
        decision = str(raw).strip().upper() if raw else ""

        if not decision:
            buckets["non_traitees"].append(item)
        elif decision not in VALID:
            item["decision_lue"] = decision
            buckets["invalides"].append(item)
        elif decision == "OUI":
            buckets["appliquer"].append(item)
        elif decision == "NON":
            buckets["abandonner"].append(item)
        elif decision == "PLUS TARD":
            buckets["parking"].append(item)
        else:
            buckets["a_modifier"].append(item)

    print(json.dumps(buckets, ensure_ascii=False, indent=2))

    n = {k: len(v) for k, v in buckets.items()}
    print(
        f"\n--- {n['appliquer']} à appliquer · {n['a_modifier']} à retravailler · "
        f"{n['parking']} en parking · {n['abandonner']} abandonnées · "
        f"{n['non_traitees']} NON TRAITÉES · {n['invalides']} invalides ---",
        file=sys.stderr,
    )
    if n["non_traitees"] or n["invalides"]:
        print(
            "⚠️  Des lignes sont sans décision valide : NE PAS les appliquer, "
            "les redemander à Adrian.",
            file=sys.stderr,
        )


if __name__ == "__main__":
    main()
